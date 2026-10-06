"""`qualify config import`: внутренние файлы компании -> локальные конфиги.

  * таблица CPA по странам (xlsx)          -> cpa_by_country.csv
  * ROI-шаблон (xlsx)                       -> volume_criteria.yaml (тиры, комиссии)
  * регламент whitelisting (docx)           -> criteria.yaml (критерий соцсетей)
                                              + volume_criteria.yaml (правила
                                              whitelisting, лимиты приглашённых)

Результат — только локальные файлы из .gitignore: внутренние числа не попадают
в репозиторий и не уходят в LLM. Ячейки и таблицы ищутся по подписям, а не по
адресам, чтобы новая версия файла не ломала импорт.
"""

from __future__ import annotations

import csv
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Any, Optional
from xml.etree import ElementTree as ET

import yaml

from .config import FEE_KEYS, ConfigError


class ImportProblem(ConfigError):
    """Файл не похож на ожидаемый — сообщение показывается как есть."""


def _openpyxl():
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - зависит от окружения
        raise ImportProblem("Для импорта xlsx нужен openpyxl: pip install openpyxl") from exc
    return openpyxl


def _norm(value: Any) -> str:
    return " ".join(str(value or "").split()).lower()


def parse_amount(text: Any) -> Optional[float]:
    """«0.2M», «500M», «2.5m», «$1,000,000», «50k» -> число."""
    if text is None:
        return None
    if isinstance(text, (int, float)):
        return float(text)
    match = re.search(r"\$?\s*([\d.,]+)\s*([kmb])?", str(text).strip(), re.IGNORECASE)
    if not match:
        return None
    number = float(match.group(1).replace(",", ""))
    mult = {"k": 1e3, "m": 1e6, "b": 1e9}.get((match.group(2) or "").lower(), 1)
    return number * mult


# --------------------------------------------------------------------------
# CPA по странам
# --------------------------------------------------------------------------

@dataclass
class CpaImport:
    rows: list[dict[str, Any]]
    sheet: str
    column: str
    skipped: int = 0


def _pick_sheet(names: list[str]) -> str:
    """Свежий квартал: «Q3 CP1T» важнее «Q2 CP1T»; год, если указан, — ещё важнее."""
    def key(name: str) -> tuple[int, int]:
        year = re.search(r"(20\d\d)", name)
        quarter = re.search(r"Q([1-4])", name, re.IGNORECASE)
        return (int(year.group(1)) if year else 0, int(quarter.group(1)) if quarter else 0)

    return max(names, key=key)


def import_cpa(path: str | Path, sheet: Optional[str] = None, column: Optional[str] = None) -> CpaImport:
    wb = _openpyxl().load_workbook(path, data_only=True, read_only=True)
    name = sheet or _pick_sheet(wb.sheetnames)
    if name not in wb.sheetnames:
        raise ImportProblem(f"{path}: нет листа «{name}». Есть: {', '.join(wb.sheetnames)}")
    rows = [list(r) for r in wb[name].iter_rows(values_only=True)]

    header_idx = next(
        (i for i, row in enumerate(rows) if any(_norm(c) == "country" for c in row)), None
    )
    if header_idx is None:
        raise ImportProblem(f"{path} / {name}: не найдена строка заголовка с «Country»")
    header = [_norm(c) for c in rows[header_idx]]
    country_col = header.index("country")
    region_col = next((i for i, h in enumerate(header) if h == "region"), None)
    if column:
        wanted = _norm(column)
        cpa_col = next((i for i, h in enumerate(header) if h == wanted), None)
        if cpa_col is None:
            raise ImportProblem(f"{path} / {name}: нет колонки «{column}». Есть: {[c for c in rows[header_idx] if c]}")
    else:
        # ставка для партнёров: «CPA Referral Pro (Affiliates)», «Affiliates CPA»
        cpa_col = next((i for i, h in enumerate(header) if "affiliate" in h), None)
        if cpa_col is None:
            raise ImportProblem(
                f"{path} / {name}: не найдена колонка ставки для партнёров (со словом Affiliate). "
                "Укажите её явно: --cpa-column"
            )
    above = " ".join(_norm(c) for r in rows[:header_idx] for c in r if c)
    event = "FTT"  # CP1T — cost per first trade
    del above

    result: list[dict[str, Any]] = []
    skipped = 0
    for row in rows[header_idx + 1:]:
        code = str(row[country_col] or "").strip().upper() if country_col < len(row) else ""
        value = row[cpa_col] if cpa_col < len(row) else None
        if not re.fullmatch(r"[A-Z]{2}", code):
            continue
        amount = parse_amount(value)
        if amount is None:
            skipped += 1
            continue
        result.append({
            "country_code": code,
            "country": "",
            "region": str(row[region_col] or "").strip() if region_col is not None else "",
            "cpa": round(amount, 2),
            "currency": "USD",
            "event": event,
        })
    if not result:
        raise ImportProblem(f"{path} / {name}: не найдено ни одной ставки")
    return CpaImport(rows=result, sheet=name, column=str(rows[header_idx][cpa_col]), skipped=skipped)


# --------------------------------------------------------------------------
# ROI-шаблон
# --------------------------------------------------------------------------

def _cells(ws: Any) -> dict[tuple[int, int], Any]:
    return {(c.row, c.column): c.value for row in ws.iter_rows() for c in row if c.value is not None}


def _find(cells: dict[tuple[int, int], Any], *labels: str) -> list[tuple[int, int]]:
    wanted = {_norm(label) for label in labels}
    return sorted(pos for pos, value in cells.items() if _norm(value) in wanted)


def _right_number(cells: dict[tuple[int, int], Any], pos: tuple[int, int]) -> Optional[float]:
    row, col = pos
    for offset in range(1, 4):
        value = cells.get((row, col + offset))
        if isinstance(value, (int, float)):
            return float(value)
    return None


def import_roi(path: str | Path) -> dict[str, Any]:
    wb = _openpyxl().load_workbook(path, data_only=False)
    ws = wb.worksheets[0]
    cells = _cells(ws)

    months = 3
    for value in cells.values():
        match = re.search(r"(\d+)\s*month evaluation", _norm(value))
        if match:
            months = int(match.group(1))
            break

    spot_hdr = _find(cells, "Spot Tier")
    fut_hdr = _find(cells, "Futures Tier")
    if not spot_hdr or not fut_hdr:
        raise ImportProblem(f"{path}: не найдена таблица тиров (заголовки «Spot Tier» / «Futures Tier»)")

    def tiers(header: tuple[int, int]) -> list[dict[str, float]]:
        row, col = header
        result = []
        r = row + 1
        while isinstance(cells.get((r, col)), (int, float)):
            rate, volume, traders = cells.get((r, col)), cells.get((r, col + 1)), cells.get((r, col + 2))
            result.append({
                "rate": round(float(rate) * 100 if float(rate) <= 1 else float(rate), 2),
                "volume": float(parse_amount(volume) or 0),
                "new_traders": int(float(traders or 0)),
            })
            r += 1
        if not result:
            raise ImportProblem(f"{path}: под «{cells[header]}» нет строк тиров")
        return result

    def default_rate(label: str) -> float:
        positions = _find(cells, label)
        for pos in positions:
            formula = str(cells.get((pos[0], pos[1] + 1)) or "")
            match = re.search(r",\s*([0-9.]+)\s*\)+\s*$", formula)
            if match:
                value = float(match.group(1))
                return round(value * 100 if value <= 1 else value, 2)
        return 0.0

    def fee_pair(*labels: str) -> tuple[float, Optional[float]]:
        positions = _find(cells, *labels)
        values = [(pos, _right_number(cells, pos)) for pos in positions]
        values = [(pos, v) for pos, v in values if v is not None]
        if not values:
            raise ImportProblem(f"{path}: не найдено значение «{labels[0]}»")
        values.sort(key=lambda pv: pv[0][1])  # левее — наши, правее — конкурента
        ours = values[0][1]
        competitor = values[-1][1] if len(values) > 1 else None
        return ours, competitor

    fee_labels = {
        "spot_taker": ("Spot Taker Rate",),
        "spot_maker": ("Spot Maker Rate", "Spot Mater Rate"),
        "futures_taker": ("Futures Taker Rate", "Future Taker Rate"),
        "futures_maker": ("Futures Maker Rate", "Future Maker Rate"),
    }
    ours: dict[str, float] = {}
    competitor: dict[str, float] = {}
    for key in FEE_KEYS:
        o, c = fee_pair(*fee_labels[key])
        ours[key] = o
        if c is not None:
            competitor[key] = c
    if len(competitor) != len(FEE_KEYS):
        competitor = dict(ours)

    def share(label: str) -> float:
        positions = _find(cells, label)
        value = _right_number(cells, positions[0]) if positions else None
        if value is None:
            raise ImportProblem(f"{path}: не найдено «{label}»")
        return value

    return {
        "evaluation_months": months,
        "markets": {
            "spot": {"default_rate": default_rate("Spot Reached Tier"), "tiers": tiers(spot_hdr[0])},
            "futures": {"default_rate": default_rate("Futures Reached Tier"), "tiers": tiers(fut_hdr[0])},
        },
        "taker_share": {
            "spot": share("Avg Spot Taker Volume Percentage"),
            "futures": share("Avg Futures Taker Volume Percentage"),
        },
        "fees": {"ours": ours, "competitor": competitor},
    }


# --------------------------------------------------------------------------
# Регламент whitelisting (docx)
# --------------------------------------------------------------------------

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _docx_tables(path: str | Path) -> list[list[list[str]]]:
    with zipfile.ZipFile(path) as z:
        root = ET.fromstring(z.read("word/document.xml"))
    tables = []
    for tbl in root.iter(_W + "tbl"):
        rows = []
        for tr in tbl.findall(_W + "tr"):
            cells = []
            for tc in tr.findall(_W + "tc"):
                parts = []
                for p in tc.iter(_W + "p"):
                    parts.append("".join(t.text or "" for t in p.iter(_W + "t")))
                cells.append(" ".join(x for x in parts if x).strip())
            rows.append(cells)
        tables.append(rows)
    return tables


_FOLLOWERS_RE = re.compile(r"([\d,]+)\s*\+\s*followers or subscribers", re.IGNORECASE)
_COMMUNITY_RE = re.compile(r"community of\s*([\d,]+)\s*\+\s*members", re.IGNORECASE)
_COMBO_RE = re.compile(r"([\d.,]+\s*[kKmM]?)\s*\+\s*followers with\s*([\d.,]+\s*[kKmM]?)\s*average views", re.IGNORECASE)
_ACHIEVE_RE = re.compile(
    r"(\d+)\s*%\s*achievement of (trading volume|new traders?)\s*criteria(?:\s*and\s*(\d+)\s*%)?",
    re.IGNORECASE,
)
_MONTHS_RE = re.compile(r"up to\s*(\d+)\s*months?", re.IGNORECASE)


@dataclass
class GuidelinesImport:
    social: dict[str, Any] = field(default_factory=dict)
    eligibility: list[dict[str, Any]] = field(default_factory=list)
    invitee_limits: dict[str, str] = field(default_factory=dict)
    volume_tiers: dict[str, list[dict[str, float]]] = field(default_factory=dict)


def import_guidelines(path: str | Path) -> GuidelinesImport:
    result = GuidelinesImport()
    tables = _docx_tables(path)

    for table in tables:
        flat = " ".join(" ".join(r) for r in table).lower()

        # --- автооценка: тиры Spot/Futures и лимиты приглашённых ---------------
        if "invitee" in flat and "tier" in flat:
            spot: list[dict[str, float]] = []
            futures: list[dict[str, float]] = []
            for row in table:
                label = _norm(row[0]) if row else ""
                if label.startswith("tier") and len(row) >= 8:
                    def tier(nt: str, vol: str, rate: str) -> dict[str, float]:
                        return {"rate": float(rate.strip().rstrip("%")), "volume": parse_amount(vol) or 0.0,
                                "new_traders": int(parse_amount(nt) or 0)}
                    spot.append(tier(row[1], row[2], row[3]))
                    futures.append(tier(row[5], row[6], row[7]))
                if label.startswith("invitee"):
                    limits = [c for c in row[1:] if c.strip()]
                    if limits:
                        result.invitee_limits["spot"] = limits[0]
                        result.invitee_limits["futures"] = limits[-1]
            if spot:
                result.volume_tiers = {"spot": spot, "futures": futures}
            continue

        # --- Criteria 1: правила whitelisting по объёмам -----------------------
        if "achievement" in flat and "eligib" in flat:
            months = None
            seen: set[tuple] = set()
            for row in table[1:]:
                text = " ".join(row)
                month_match = _MONTHS_RE.search(row[-1] if row else "")
                if month_match:
                    months = int(month_match.group(1))
                achieve = _ACHIEVE_RE.search(text)
                if not achieve or months is None:
                    continue
                top = "top" in text.lower() and "ftt" in text.lower()
                rule = {
                    "primary_pct": float(achieve.group(1)),
                    "secondary_pct": 0.0 if top else float(achieve.group(3) or 0),
                    "months": months,
                }
                if top:
                    rule["requires_top_ftt"] = True
                key = tuple(sorted(rule.items()))
                if key not in seen:
                    seen.add(key)
                    result.eligibility.append(rule)
            continue

        # --- Criteria 2: соцсети ----------------------------------------------
        if "followers or subscribers" in flat:
            types: dict[str, dict[float, dict[str, Any]]] = {"individual": {}, "institutional": {}}
            rate: Optional[float] = None
            for row in table[1:]:
                if len(row) < 3:
                    continue
                rate_match = re.search(r"(\d+(?:\.\d+)?)\s*%", row[0])
                if rate_match:
                    rate = float(rate_match.group(1))
                if rate is None:
                    continue
                kind = "institutional" if "institutional" in row[1].lower() else "individual"
                profile = row[2]
                tier: dict[str, Any] = {"rate": rate}
                followers = _FOLLOWERS_RE.search(profile)
                if followers:
                    tier["followers_single_platform"] = int(followers.group(1).replace(",", ""))
                community = _COMMUNITY_RE.search(profile)
                if community:
                    tier["community_members"] = int(community.group(1).replace(",", ""))
                combo = _COMBO_RE.search(profile)
                if combo:
                    tier["followers_with_views"] = {
                        "followers": int(parse_amount(combo.group(1)) or 0),
                        "avg_views": int(parse_amount(combo.group(2)) or 0),
                    }
                if len(tier) > 1:
                    types[kind][rate] = tier
            result.social = {
                kind: {"match": "any", "tiers": sorted(tiers.values(), key=lambda t: -t["rate"])}
                for kind, tiers in types.items() if tiers
            }
    return result


# --------------------------------------------------------------------------
# Запись конфигов
# --------------------------------------------------------------------------

def _header(sources: list[str]) -> str:
    return (
        f"# Сгенерировано: qualify config import, {date.today().isoformat()}\n"
        f"# Источники: {', '.join(sources)}\n"
        "# ВНУТРЕННИЕ ДАННЫЕ: файл в .gitignore, не коммитить и не пересылать.\n\n"
    )


def build_volume_yaml(roi: Optional[dict[str, Any]], guidelines: Optional[GuidelinesImport],
                      existing: Optional[dict[str, Any]] = None) -> dict[str, Any]:
    data: dict[str, Any] = dict(existing or {})
    if roi:
        data.update({k: v for k, v in roi.items() if k != "markets"})
        markets = data.setdefault("markets", {})
        for market, raw in roi["markets"].items():
            markets.setdefault(market, {}).update(raw)
    if guidelines:
        markets = data.setdefault("markets", {})
        for market, tiers in guidelines.volume_tiers.items():
            if not (markets.get(market) or {}).get("tiers"):
                markets.setdefault(market, {})["tiers"] = tiers
        for market, limit in guidelines.invitee_limits.items():
            markets.setdefault(market, {})["invitee_limit"] = limit
        if guidelines.eligibility:
            data["eligibility"] = guidelines.eligibility
    return data


def tier_mismatches(roi: dict[str, Any], guidelines: GuidelinesImport) -> list[str]:
    """Пороги тиров в ROI-шаблоне и в регламенте должны совпадать."""
    problems = []
    for market, doc_tiers in guidelines.volume_tiers.items():
        roi_tiers = {t["rate"]: t for t in roi["markets"][market]["tiers"]}
        for tier in doc_tiers:
            other = roi_tiers.get(tier["rate"])
            if other and (abs(other["volume"] - tier["volume"]) > 1 or other["new_traders"] != tier["new_traders"]):
                problems.append(
                    f"{market} {tier['rate']:g}%: в ROI-шаблоне {other['volume']:,.0f}$ / {other['new_traders']}, "
                    f"в регламенте {tier['volume']:,.0f}$ / {tier['new_traders']} — взяты значения ROI-шаблона"
                )
    return problems


def write_cpa(path: Path, rows: list[dict[str, Any]], sources: list[str]) -> None:
    with path.open("w", encoding="utf-8", newline="") as fh:
        fh.write(_header(sources))
        writer = csv.DictWriter(fh, fieldnames=["country_code", "country", "region", "cpa", "currency", "event"])
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def write_yaml(path: Path, data: dict[str, Any], sources: list[str]) -> None:
    text = yaml.safe_dump(data, allow_unicode=True, sort_keys=False, default_flow_style=None, width=100)
    path.write_text(_header(sources) + text, encoding="utf-8")


def build_criteria_yaml(social: dict[str, Any], existing: Optional[dict[str, Any]] = None) -> dict[str, Any]:
    data: dict[str, Any] = dict(existing or {})
    data["affiliate_types"] = social
    data.setdefault("cpa_rates_file", "cpa_by_country.csv")
    data.setdefault("borderline_margin", 0.10)
    data.setdefault("deal", {
        "test_period_days": 30,
        "review_metric": "share of registered users who started trading, and trading volume",
    })
    return data
